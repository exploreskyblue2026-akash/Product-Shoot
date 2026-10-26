"""Product Library (v2, phase 1).

Everything dumped in (master sheet, shoot sheet, photos, ZIPs, catalogue PDFs, a Drive folder) is sorted by product code:
images are matched to codes, duplicates dropped, the best image per product marked primary, cut-outs made in the
browser and stored here (+ Drive "Cutouts library"), and a report lists what is ready and what is missing.
"""
import io
import re
import zipfile
from difflib import get_close_matches

from PIL import Image
from sqlalchemy import select, delete
from sqlalchemy.orm import Session

from .connectors import storage
from .db import LibImage, LibProduct, get_setting, set_setting, log

IMG_EXT = (".jpg", ".jpeg", ".png", ".webp")
LOW_RES = 800            # long side below this = flagged "low resolution"
SKIP_FOLDERS = ("final output", "cutouts library", "product library")


def norm(code) -> str:
    """'BCO64/GY', 'bco64-gy', 'BCO64 GY' -> 'BCO64GY'. Numbers from Excel (76618.0) -> '76618'."""
    if code is None:
        return ""
    if isinstance(code, float) and code.is_integer():
        code = int(code)
    return re.sub(r"[^A-Z0-9]", "", str(code).upper())


# ------------------------------------------------------------------ master sheet
ALIASES = {
    "product_id": ("product id", "productid", "reference", "model", "article", "art no", "sku code"),
    "item_code": ("item code", "itemcode", "item no", "code"),
    "ean": ("eancode", "ean code", "ean", "barcode", "gtin"),
    "name": ("item name", "product name", "name", "title"),
    "category": ("short name", "category", "type"),
}


def _header_map(head: list[str]) -> dict[str, int]:
    h = [str(x or "").strip().lower() for x in head]
    out = {}
    for field, names in ALIASES.items():
        for n in names:
            if n in h:
                out[field] = h.index(n)
                break
    return out


def _cell(row, i):
    return row[i] if i is not None and i < len(row) and row[i] is not None else ""


def import_master(s: Session, data: bytes, filename: str = "master.xlsx") -> dict:
    """Read every tab that has a Product ID or Item code column. Tab name = brand."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    added = updated = 0
    tabs = []
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        try:
            head = next(rows)
        except StopIteration:
            continue
        m = _header_map(list(head))
        if "product_id" not in m and "item_code" not in m:
            continue
        tabs.append(ws.title)
        for r in rows:
            pid = str(_cell(r, m.get("product_id"))).strip()
            item = norm(_cell(r, m.get("item_code")))
            key = norm(pid) or item
            if not key:
                continue
            ean = norm(_cell(r, m.get("ean")))
            p = s.scalar(select(LibProduct).where(LibProduct.key == key))
            if not p:
                p = LibProduct(key=key)
                s.add(p)
                added += 1
            else:
                updated += 1
            p.product_id, p.item_code, p.ean = pid, item, ean
            p.name = str(_cell(r, m.get("name"))).strip()[:300]
            p.category = str(_cell(r, m.get("category"))).strip()[:80]
            p.brand = ws.title.strip()[:80]
    s.commit()
    log(s, "library", f"Master sheet {filename}: {added} new, {updated} updated products from tabs {', '.join(tabs)}")
    s.commit()
    return {"added": added, "updated": updated, "tabs": tabs}


def focus_index(s: Session, idx: dict | None = None) -> dict:
    """Only the products in the shoot sheet (when one is loaded and focus is on): the scan touches nothing else."""
    idx = idx if idx is not None else _index(s)
    keys = {i["key"] for i in (get_setting(s, "shoot_list", []) or []) if i.get("key")}
    if not keys or not get_setting(s, "library_focus", True):
        return idx
    return {k: p for k, p in idx.items() if p.key in keys}


def _variants(key: str) -> set:
    """A code and its one-letter-off spellings (PPIP25RD ~ PIP25RD) for a quick 'is it on this page' test."""
    return {key} | {key[:i] + key[i + 1:] for i in range(len(key))} if len(key) >= 6 else {key}


def _index(s: Session) -> dict[str, LibProduct]:
    """Every way a product can be referred to (product id, item code, EAN) -> product."""
    idx = {}
    for p in s.scalars(select(LibProduct)):
        for k in (p.key, p.item_code, p.ean):
            if k:
                idx.setdefault(k, p)
    return idx


def resolve(code, idx: dict[str, LibProduct]) -> tuple[LibProduct | None, str]:
    k = norm(code)
    if not k:
        return None, ""
    if k in idx:
        return idx[k], "exact"
    close = get_close_matches(k, [x for x in idx if len(x) >= 5], n=1, cutoff=0.88)
    return (idx[close[0]], f"close match ({close[0]})") if close else (None, "")


# ------------------------------------------------------------------ shoot sheet
def import_shoot(s: Session, data: bytes, filename: str = "shoot.xlsx") -> dict:
    """Columns: SHOOT (code), SHOTS (e.g. 1,3,5), PRIORITY, NOTE. Codes may be product ids, item codes or EANs."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    head = [str(x or "").strip().lower() for x in next(rows)]
    col = lambda *names: next((head.index(n) for n in names if n in head), None)   # noqa: E731
    ci, si, pi, ni = col("shoot", "code", "product code", "product id", "item code"), col("shots"), col("priority"), col("note", "notes")
    if ci is None:
        ci = 0
    idx, items = _index(s), []
    for n, r in enumerate(rows):
        raw = _cell(r, ci)
        if raw in ("", None):
            continue
        p, how = resolve(raw, idx)
        shots = [int(x) for x in re.findall(r"[1-5]", str(_cell(r, si)))] or []
        try:
            prio = int(float(_cell(r, pi) or 0)) or (1000 + n)
        except ValueError:
            prio = 1000 + n
        items.append({"code": str(raw).strip() if not isinstance(raw, float) else str(int(raw)),
                      "key": p.key if p else "", "match": how or "not in master sheet",
                      "shots": sorted(set(shots)), "priority": prio, "note": str(_cell(r, ni)).strip()})
    items.sort(key=lambda x: x["priority"])
    set_setting(s, "shoot_list", items)
    log(s, "library", f"Shoot sheet {filename}: {len(items)} products, {sum(1 for i in items if i['key'])} found in master")
    s.commit()
    return {"items": len(items), "found": sum(1 for i in items if i["key"])}


# ------------------------------------------------------------------ images
def dhash(im: Image.Image) -> str:
    g = im.convert("L").resize((9, 8), Image.LANCZOS)
    px = list(g.getdata())
    bits = 0
    for r in range(8):
        for c in range(8):
            bits = (bits << 1) | (px[r * 9 + c] > px[r * 9 + c + 1])
    return f"{bits:016x}"


def _ham(a: str, b: str) -> int:
    return bin(int(a, 16) ^ int(b, 16)).count("1") if a and b else 64


def key_from_name(name: str, idx: dict[str, LibProduct]) -> LibProduct | None:
    """Find a known code inside a file or folder name: '6110224MET_front.jpg', 'BCO64-GY 2.png', '76618/IMG_1.jpg'."""
    parts = [norm(x) for x in re.split(r"[\\/]", name)]
    stem_parts = [norm(re.sub(r"\.[a-z0-9]{3,4}$", "", x, flags=re.I)) for x in re.split(r"[\\/]", name)]
    best = None
    for part in stem_parts[::-1] + parts[::-1]:
        if not part:
            continue
        if part in idx:
            return idx[part]
        for k in idx:                                    # longest code contained in the name wins
            if len(k) >= 5 and k in part and (best is None or len(k) > len(best[0])):
                best = (k, idx[k])
        # tokens split on _ - space, e.g. "TRV20-GY_1"
    if best:
        return best[1]
    toks = re.split(r"[_\s.]+", re.sub(r"\.[a-z0-9]{3,4}$", "", name.split("/")[-1], flags=re.I))
    for i in range(len(toks), 0, -1):
        k = norm("".join(toks[:i]))
        if k in idx:
            return idx[k]
    return None


def add_image(s: Session, data: bytes, name: str, source: str = "file", key: str | None = None,
              drive_id: str = "", idx=None) -> LibImage | None:
    try:
        im = Image.open(io.BytesIO(data))
        im.load()
    except Exception:
        return None
    if min(im.size) < 120:                               # icons, logos, bullets in catalogues
        return None
    idx = idx if idx is not None else _index(s)
    if key is None:
        p = key_from_name(name, idx)
        key = p.key if p else ""
    h = dhash(im)
    rel = storage.save(f"library/raw/{storage.safe(key or 'unmatched')}/{storage.safe(name)[-120:]}",
                       storage.normalise_raw(data))
    img = LibImage(key=key, source=source, src_name=name[:300], drive_id=drive_id, path=rel,
                   w=im.size[0], h=im.size[1], phash=h)
    for other in s.scalars(select(LibImage).where(LibImage.key == key, LibImage.dup_of.is_(None))):
        if _ham(other.phash, h) <= 5:                    # same picture (even resized / re-saved)
            if img.w * img.h > other.w * other.h:        # keep the bigger one as the original
                other.dup_of = None
                s.add(img)
                s.flush()
                other.dup_of = img.id
            else:
                img.dup_of = other.id
            break
    s.add(img)
    s.flush()
    if key:
        pick_primary(s, key)
        if img.dup_of is None and not cutouts_on(s) and get_setting(s, "library_drive_copy", False):
            save_original_to_drive(s, img, data)
    return img


def pick_primary(s: Session, key: str, force_id: int | None = None):
    """Biggest image wins, unless you chose one or the current primary already has its cut-out."""
    imgs = list(s.scalars(select(LibImage).where(LibImage.key == key, LibImage.dup_of.is_(None))))
    if not imgs:
        return
    cur = next((i for i in imgs if i.primary), None)
    best = next((i for i in imgs if i.id == force_id), None)
    if best is None:
        usable = [i for i in imgs if i.cut_status != "skip"] or imgs
        best = cur if cur and cur.cut_status in ("ok", "check") else max(usable, key=lambda i: i.w * i.h)
    for i in imgs:
        i.primary = i is best
    for i in s.scalars(select(LibImage).where(LibImage.key == key, LibImage.dup_of.is_not(None))):
        i.primary = False


def add_files(s: Session, files: list[tuple[str, bytes]]) -> dict:
    """Uploaded files: images, ZIPs, catalogue PDFs, master/shoot sheets (detected by their columns)."""
    stats = {"images": 0, "catalogue_images": 0, "sheets": [], "skipped": 0}
    idx = _index(s)
    for name, data in files:
        low = name.lower()
        if low.endswith((".xlsx", ".xlsm")):
            stats["sheets"].append(add_sheet(s, data, name))
            idx = _index(s)
        elif low.endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for n in z.namelist():
                    if n.lower().endswith(IMG_EXT) and not n.startswith("__MACOSX"):
                        stats["images"] += bool(add_image(s, z.read(n), n, "file", idx=idx))
        elif low.endswith(".pdf"):
            stats["catalogue_images"] += add_catalogue(s, data, name, idx=focus_index(s, idx))
        elif low.endswith(IMG_EXT):
            stats["images"] += bool(add_image(s, data, name, "file", idx=idx))
        else:
            stats["skipped"] += 1
    s.commit()
    return stats


def add_sheet(s: Session, data: bytes, name: str) -> str:
    """Decide whether an .xlsx is the master sheet or the shoot sheet from its first row."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    head = [str(x or "").strip().lower() for x in next(wb.worksheets[0].iter_rows(values_only=True), [])]
    if "shoot" in head or ("shots" in head and len(head) <= 6):
        r = import_shoot(s, data, name)
        return f"{name}: shoot sheet, {r['items']} products ({r['found']} found)"
    r = import_master(s, data, name)
    return f"{name}: master sheet, {r['added']} new + {r['updated']} updated products"


# ------------------------------------------------------------------ catalogue PDF
CAPTION = re.compile(r"[|#]\s*#?\s*([A-Z0-9]{2,}[A-Z0-9/\-]*[A-Z0-9])", re.I)


def add_catalogue(s: Session, pdf: bytes, filename: str, idx=None) -> int:
    """Each product picture has its caption ('NAME | CODE') just below it. Match by position; unmatched pictures
    go to the Unmatched tray. Full-width banners and backgrounds are ignored."""
    import fitz
    idx = idx if idx is not None else _index(s)
    doc = fitz.open(stream=pdf, filetype="pdf")
    added = 0
    want = set().union(*[_variants(p.key) for p in set(idx.values())]) if idx else set()
    for pno, page in enumerate(doc, 1):
        if want and len(idx) < 400:                    # focus mode: skip pages that mention none of the codes
            flat = norm(page.get_text())
            if not any(v in flat for v in want):
                continue
        W = page.rect.width
        pics = [i for i in page.get_image_info(xrefs=True)
                if i["xref"] and (i["bbox"][2] - i["bbox"][0]) < W * 0.7 and (i["bbox"][3] - i["bbox"][1]) > 40]
        caps = []
        for b in page.get_text("blocks"):
            text = b[4]
            codes = [c for c in CAPTION.findall(text)]
            codes += [t for t in re.findall(r"\b[A-Z0-9]{3,}(?:/[A-Z0-9]{1,4})?\b", text) if norm(t) in idx]
            for c in codes:
                p = idx.get(norm(c)) or idx.get(norm(re.sub(r"-[A-Z0-9]+$", "", c)))   # '#TRV20/BK-A037'
                if not p and "/" in c:
                    p = resolve(c, idx)[0]                     # PIP25/RD in the catalogue = PPIP25/RD in the sheet
                if p:
                    caps.append((b[:4], p))
                    break
        used = set()
        for (x0, y0, x1, y1), p in caps:
            best, gap = None, 1e9
            for n, im in enumerate(pics):
                bx0, by0, bx1, by1 = im["bbox"]
                overlap = min(x1, bx1) - max(x0, bx0)
                g = y0 - by1                                   # caption top minus picture bottom
                if overlap > 10 and -15 <= g < 120 and g < gap and n not in used:
                    best, gap = n, g
            if best is None:
                continue
            used.add(best)
            pix = fitz.Pixmap(doc, pics[best]["xref"])
            if pix.n - pix.alpha >= 4:
                pix = fitz.Pixmap(fitz.csRGB, pix)
            if add_image(s, pix.tobytes("png"), f"{filename} p{pno} {p.product_id or p.key}.png", "catalogue",
                         key=p.key, idx=idx):
                added += 1
        # pictures on a page with product codes that could not be placed: Unmatched tray, with the page's codes as hints
        if caps:
            hint = ", ".join(dict.fromkeys(p.product_id or p.key for _, p in caps))
            for n, im in enumerate(pics):
                if n in used or im["width"] < 200 or im["height"] < 150:
                    continue
                pix = fitz.Pixmap(doc, im["xref"])
                if pix.n - pix.alpha >= 4:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                add_image(s, pix.tobytes("png"), f"{filename} p{pno} (codes on page: {hint}) #{n}.png", "catalogue",
                          key="", idx=idx)
    log(s, "library", f"Catalogue {filename}: {added} product pictures matched")
    return added


# ------------------------------------------------------------------ Drive
def scan_drive(s: Session, link: str) -> dict:
    """Scan a Drive folder (or the whole Shared drive): sheets, PDFs and images. Final output / Cutouts library skipped."""
    from .connectors import google
    svc, root = google.drive(), google.folder_id(link)
    known = {i.drive_id for i in s.scalars(select(LibImage).where(LibImage.drive_id != ""))}
    files, stack = [], [(root, "")]
    while stack:
        cur, prefix = stack.pop()
        token = None
        while True:
            r = svc.files().list(q=f"'{cur}' in parents and trashed=false", pageToken=token, pageSize=200,
                                 fields="nextPageToken, files(id,name,mimeType,md5Checksum,modifiedTime)", supportsAllDrives=True,
                                 includeItemsFromAllDrives=True, corpora="allDrives").execute()
            for f in r.get("files", []):
                if f["mimeType"] == "application/vnd.google-apps.folder":
                    if f["name"].strip().lower() == "cutouts library":
                        set_setting(s, "cutouts_drive_folder", f["id"])
                    if f["name"].strip().lower() == "product library":
                        set_setting(s, "library_drive_folder", f["id"])
                    if f["name"].strip().lower() not in SKIP_FOLDERS:
                        stack.append((f["id"], f"{prefix}{f['name']}/"))
                else:
                    files.append({**f, "path": prefix + f["name"]})
            token = r.get("nextPageToken")
            if not token:
                break
    sheets = [f for f in files if f["name"].lower().endswith((".xlsx", ".xlsm"))
              or f["mimeType"] == "application/vnd.google-apps.spreadsheet"]
    stats = {"sheets": [], "images": 0, "catalogue_images": 0, "already": 0}
    # sheets first (master before shoot, so codes can be resolved)
    loaded = []
    seen = dict(get_setting(s, "sheet_versions", {}) or {})
    for f in sheets:
        ver = f.get("md5Checksum") or f.get("modifiedTime") or ""
        if ver and seen.get(f["id"]) == ver:            # unchanged since the last scan: skip
            continue
        seen[f["id"]] = ver
        if f["mimeType"] == "application/vnd.google-apps.spreadsheet":
            data = svc.files().export(fileId=f["id"], mimeType=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")).execute()
        else:
            data = google.download(f["id"])
        loaded.append((f["name"], data))
    loaded.sort(key=lambda x: 0 if _is_master(x[1]) else 1)
    for name, data in loaded:
        stats["sheets"].append(add_sheet(s, data, name))
    set_setting(s, "sheet_versions", seen)
    s.commit()
    idx = focus_index(s)
    focused = len(idx) < len(_index(s))
    stats["skipped_not_in_sheet"] = 0
    for f in files:
        if f["id"] in known:
            stats["already"] += 1
            continue
        low = f["name"].lower()
        if low.endswith(".pdf"):
            stats["catalogue_images"] += add_catalogue(s, google.download(f["id"]), f["name"], idx=idx)
            s.add(LibImage(key="", source="marker", drive_id=f["id"], cut_status="skip", src_name=f["name"]))
        elif f["mimeType"].startswith("image/") or low.endswith(IMG_EXT):
            p = key_from_name(f["path"], idx)
            if focused and not p:                      # not a shoot-sheet product: don't even download it
                stats["skipped_not_in_sheet"] += 1
                continue
            if add_image(s, google.download(f["id"]), f["path"], "drive", key=p.key if p else "", drive_id=f["id"], idx=idx):
                stats["images"] += 1
        s.commit()
    set_setting(s, "library_drive_link", link)
    log(s, "library", f"Drive scan: {stats['images']} photos, {stats['catalogue_images']} catalogue pictures, "
                      f"{len(stats['sheets'])} sheets")
    s.commit()
    return stats


def _is_master(data: bytes) -> bool:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        head = [str(x or "").strip().lower() for x in next(wb.worksheets[0].iter_rows(values_only=True), [])]
        return "shoot" not in head
    except Exception:
        return False


# ------------------------------------------------------------------ cut-outs
def save_cutout(s: Session, img: LibImage, png: bytes) -> LibImage:
    """PNG with transparency from the browser. Flag weak results for a person to check."""
    im = Image.open(io.BytesIO(png)).convert("RGBA")
    a = im.split()[-1]
    hist = a.histogram()
    total = im.size[0] * im.size[1]
    solid = sum(hist[200:]) / total
    soft = sum(hist[20:200]) / total
    bbox = a.point(lambda v: 255 if v > 20 else 0).getbbox()
    notes = []
    if solid < 0.03:
        notes.append("almost nothing kept")
    if soft > 0.08:
        notes.append("soft or see-through edges")
    if bbox and (bbox[0] <= 1 or bbox[1] <= 1 or bbox[2] >= im.size[0] - 1 or bbox[3] >= im.size[1] - 1):
        notes.append("product touches the photo edge, may be cropped")
    if max(img.w, img.h) < LOW_RES:
        notes.append(f"low resolution ({img.w}×{img.h})")
    if bbox:
        im = im.crop(bbox)
    out = io.BytesIO()
    im.save(out, "PNG", optimize=True)
    data = out.getvalue()
    img.cut_path = storage.save(f"library/cutouts/{storage.safe(img.key or 'unmatched')}/{img.id}.png", data)
    img.cut_status = "check" if notes else "ok"
    img.cut_note = "; ".join(notes)
    folder = get_setting(s, "cutouts_drive_folder", "")
    if folder and img.key:
        try:
            from .connectors import google
            sub = google.ensure_subfolder(folder, img.key)
            google.upload(sub, f"{img.key}_{img.id}.png", data, "image/png")
        except Exception as e:                            # keep the local copy; report the Drive problem
            img.cut_note = (img.cut_note + "; " if img.cut_note else "") + f"Drive copy failed: {str(e)[:120]}"
    return img


def cutouts_on(s: Session) -> bool:
    """Cut-outs are only needed for the composite engine. The default keeps original photos untouched."""
    return bool(get_setting(s, "library_cutouts", False))


def pending(s: Session, limit: int = 50) -> list[LibImage]:
    """Primary images of matched products that still need a cut-out (shoot-list products first)."""
    if not cutouts_on(s):
        return []
    order = {i["key"]: n for n, i in enumerate(get_setting(s, "shoot_list", []) or []) if i.get("key")}
    rows = list(s.scalars(select(LibImage).where(LibImage.primary == True, LibImage.cut_status == "pending",  # noqa: E712
                                                 LibImage.key != "", LibImage.dup_of.is_(None))))
    rows.sort(key=lambda i: (order.get(i.key, 10_000), i.id))
    return rows[:limit]


# ------------------------------------------------------------------ report
def report(s: Session) -> dict:
    prods = {p.key: p for p in s.scalars(select(LibProduct))}
    imgs = list(s.scalars(select(LibImage).where(LibImage.source != "marker")))
    by = {}
    for i in imgs:
        by.setdefault(i.key, []).append(i)
    shoot = get_setting(s, "shoot_list", []) or []
    cut = cutouts_on(s)
    keys = [i["key"] for i in shoot if i.get("key")] + sorted(k for k in by if k and k not in {i.get("key") for i in shoot})
    rows = []
    for item in shoot:
        if not item.get("key"):
            rows.append({"code": item["code"], "key": "", "name": "", "brand": "", "in_shoot": True,
                         "status": "not_in_master", "images": 0, "dups": 0, "primary": None,
                         "notes": "Code not found in the master sheet", "shots": item.get("shots", []),
                         "note": item.get("note", ""), "priority": item.get("priority")})
    for k in dict.fromkeys(keys):
        p, lst = prods.get(k), by.get(k, [])
        live = [i for i in lst if i.dup_of is None]
        prim = next((i for i in live if i.primary), None)
        sh = next((i for i in shoot if i.get("key") == k), None)
        if not live:
            status, notes = "needs_photo", "No photo found: add one to Raw input named with the product id"
        elif not cut and prim and prim.cut_status != "skip":
            status, notes = "ready", ""
        elif prim and prim.cut_status == "ok":
            status, notes = "ready", ""
        elif prim and prim.cut_status == "check":
            status, notes = "check", prim.cut_note
        elif prim and prim.cut_status == "skip":
            status, notes = "skipped", "Skipped by you"
        else:
            status, notes = "cutting", "Waiting for cut-out (keep the Library page open)"
        if prim and max(prim.w, prim.h) < LOW_RES and "low resolution" not in notes:
            notes = (notes + "; " if notes else "") + f"low resolution ({prim.w}×{prim.h})"
        rows.append({"code": (p.product_id or k) if p else k, "key": k, "name": p.name if p else "",
                     "brand": p.brand if p else "", "category": p.category if p else "", "in_shoot": bool(sh),
                     "status": status, "images": len(live), "dups": len(lst) - len(live),
                     "primary": _img(prim) if prim else None, "notes": notes,
                     "shots": (sh or {}).get("shots", []), "note": (sh or {}).get("note", ""),
                     "priority": (sh or {}).get("priority"), "match": (sh or {}).get("match", ""),
                     "all": [_img(i) for i in live]})
    unmatched = [_img(i) for i in by.get("", []) if i.dup_of is None]
    counts = {}
    for r in rows:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    return {"rows": rows, "unmatched": unmatched, "counts": counts, "products": len(prods),
            "images": len([i for i in imgs if i.dup_of is None]), "shoot": len(shoot), "cutouts": cut}


def refs_for(s: Session, key: str, limit: int = 3) -> list[str]:
    """Best photos of one product (primary first, then the largest others) to give the AI several angles."""
    live = [i for i in s.scalars(select(LibImage).where(LibImage.key == key, LibImage.dup_of.is_(None),
                                                         LibImage.cut_status != "skip"))]
    live.sort(key=lambda i: (not i.primary, -(i.w * i.h)))
    return [i.path for i in live[:limit]]


def save_original_to_drive(s: Session, img: LibImage, data: bytes):
    """Keep the untouched original in Drive, grouped by product code: <Product library>/<code>/<code>_<n>.<ext>."""
    folder = get_setting(s, "library_drive_folder", "") or get_setting(s, "cutouts_drive_folder", "")
    if not folder or not img.key or img.dup_of:
        return
    try:
        from .connectors import google
        p = s.scalar(select(LibProduct).where(LibProduct.key == img.key))
        code = storage.safe((p.product_id if p else img.key) or img.key)
        sub = google.ensure_subfolder(folder, code)
        ext = (re.search(r"\.(jpe?g|png|webp)$", img.src_name, re.I) or [None, "png"])[1].lower()
        name = f"{code}_{img.id}.{ext}"
        svc = google.drive()
        q = f"'{sub}' in parents and name='{name}' and trashed=false"
        if not svc.files().list(q=q, fields="files(id)", supportsAllDrives=True, includeItemsFromAllDrives=True).execute().get("files"):
            google.upload(sub, name, data, "image/png" if ext == "png" else f"image/{'jpeg' if ext.startswith('jp') else ext}")
    except Exception as e:
        img.cut_note = f"Drive copy failed: {str(e)[:120]}"


def _img(i: LibImage) -> dict:
    return {"id": i.id, "src": f"/files/{i.path}", "cut": f"/files/{i.cut_path}" if i.cut_path else "",
            "w": i.w, "h": i.h, "status": i.cut_status, "note": i.cut_note, "name": i.src_name,
            "source": i.source, "primary": i.primary}


def report_csv(s: Session) -> str:
    import csv
    r = report(s)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Code", "Name", "Brand", "In shoot sheet", "Status", "Images", "Duplicates dropped",
                "Primary size", "Shots", "Note", "Issues"])
    for x in r["rows"]:
        pr = x["primary"]
        w.writerow([x["code"], x["name"], x["brand"], "yes" if x["in_shoot"] else "", x["status"], x["images"],
                    x["dups"], f"{pr['w']}x{pr['h']}" if pr else "", ",".join(map(str, x["shots"])) or "default",
                    x["note"], x["notes"]])
    w.writerow([])
    w.writerow([f"Unmatched images: {len(r['unmatched'])}"])
    for u in r["unmatched"]:
        w.writerow(["", u["name"]])
    return buf.getvalue()


def assign(s: Session, image_id: int, code: str) -> str:
    img = s.get(LibImage, image_id)
    p, how = resolve(code, _index(s))
    if not img or not p:
        raise ValueError("Code not found in the master sheet")
    img.key = p.key
    pick_primary(s, p.key)
    if not cutouts_on(s):
        try:
            save_original_to_drive(s, img, storage.read(img.path))
        except Exception:
            pass
    s.commit()
    return p.key


def clear(s: Session):
    s.execute(delete(LibImage))
    s.commit()
